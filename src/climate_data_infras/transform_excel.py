import os 
import pandas as pd
import datetime as dt
import openpyxl as px

exist_file = [f for f in os.listdir('./data') if f.endswith(('_processed.xlsx', 'combined.xlsx'))]
for filename in exist_file:
    os.remove(os.path.join('./data',filename))

excel_files = [
    f for f in os.listdir('./data')
    if f.endswith('.xlsx')
]

for workbook in excel_files:
    file_path = os.path.join('./data', workbook)
    excel_data = px.load_workbook(file_path) 
    
    print (f"{workbook}: {excel_data.sheetnames}")

    for sheet_name in excel_data.sheetnames:
        try: 
            parse_date = dt.datetime.strptime(sheet_name, '%d%m%y')
            parse_date = parse_date.strftime('%Y-%m-%d')
            print(parse_date)
        except ValueError: 
            print(f"Skipping sheet '{sheet_name}' in workbook '{workbook}': Invalid date format.")
            continue

        sheet = excel_data[sheet_name]
        
        date_column = sheet.max_column + 1
        sheet.cell(row = 1, column = date_column).value = 'Date'

        for row in range(3, sheet.max_row + 1):
            sheet.cell(row = row, column = date_column).value = parse_date
    
    filename, extension = os.path.splitext(workbook)
    new_workbook = f"{filename}_processed{extension}"

    output_path_processed = os.path.join('./data', new_workbook)
    excel_data.save(output_path_processed)

processed_files = [
    f for f in os.listdir('./data') if f.endswith('_processed.xlsx')
]

for workbook in processed_files:
    file_path = os.path.join('./data', workbook)
    excel_data = pd.read_excel(file_path, sheet_name=None)
    try: 
        combined_data = pd.concat(excel_data.values(), ignore_index=True)
        
    except ValueError:
        print(f"Skipping workbook '{workbook}': No sheets to combine.")
        continue

    filename, extension = os.path.splitext(workbook)
    new_workbook = f"{filename}_combined{extension}"

    output_path_combined = os.path.join('./data', new_workbook)
    combined_data.to_excel(output_path_combined, index=False)